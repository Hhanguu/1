"""
fix_class_graph.py - Sua noi dung mo ta Class Graph trong docx
Cho khop voi code thuc te trong graph.py
"""

from docx import Document

INPUT = "/home/hungvo/Downloads/CauTrucRoiRacDo_Thi_Chinh_Sua.docx"
OUTPUT = "/home/hungvo/Downloads/CauTrucRoiRacDo_Thi_Chinh_Sua.docx"


def replace_in_paragraph(para, old, new):
    """Replace text in paragraph while preserving formatting."""
    if old in para.text:
        for run in para.runs:
            if old in run.text:
                run.text = run.text.replace(old, new)
                return True
    return False


def main():
    doc = Document(INPUT)
    changes = 0

    print("Sua noi dung Class Graph trong docx...\n")

    # ============================================================
    # FIX 1: Sua mo ta thuoc tinh Class Graph (Table 0 - assignment table)
    # and paragraphs describing the class
    # ============================================================

    # Fix paragraph descriptions of Class Graph attributes
    fixes = [
        # directed -> is_directed
        ("directed (bool)", "is_directed (bool)"),
        ("directed(bool)", "is_directed (bool)"),
        ("self.directed", "self.is_directed"),

        # adj_list -> adj
        ("adj_list (dict)", "adj (dict)"),
        ("adj_list(dict)", "adj (dict)"),
        ("self.adj_list", "self.adj"),
        ("adj_list là", "adj là"),
        ("adj_list]", "adj]"),
        ("adj_list}", "adj}"),
        ("adj_list=", "adj="),

        # nodes type: set -> list
        ("nodes (set)", "nodes (list)"),
        ("nodes(set)", "nodes (list)"),

        # add_node is not a separate method
        ("add_node(node) / add_edge(u, v, weight)", "add_edge(u, v, w=1)"),
        ("add_node() / add_edge", "add_edge"),
        ("add_node(node)", "add_edge(u, v, w)"),

        # get_nodes -> get_all_edges (get_nodes doesn't exist)
        ("get_nodes() / get_neighbors(node)", "get_neighbors(node) / get_all_edges()"),
        ("get_nodes()", "get_all_edges()"),

        # Method name fixes
        ("to_adjacency_matrix()", "get_adjacency_matrix()"),
        ("to_edge_list()", "get_edge_list()"),
        ("print_adj_list()", "get_adjacency_list()"),
        ("print_adj_matrix()", "get_adjacency_matrix()"),

        # draw() -> draw_graph()
        ("hàm draw()", "hàm draw_graph()"),
        ("draw()", "draw_graph()"),
        ("phương thức draw()", "phương thức draw_graph()"),

        # Other method references
        ("from_adjacency_matrix()", "get_adjacency_matrix()"),
        ("from_edge_list()", "get_edge_list()"),
    ]

    # Apply to all paragraphs
    for para in doc.paragraphs:
        for old, new in fixes:
            if old in para.text:
                if replace_in_paragraph(para, old, new):
                    changes += 1
                    print(f"  P: '{old}' -> '{new}'")

    # Apply to all tables
    for t_idx, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for old, new in fixes:
                        if old in para.text:
                            if replace_in_paragraph(para, old, new):
                                changes += 1
                                print(f"  T{t_idx}: '{old}' -> '{new}'")

    # ============================================================
    # FIX 2: Sua doan mo ta chi tiet Class Graph
    # (Table 9 - __init__, Table 10 - add_edge, etc.)
    # ============================================================

    # Fix the __init__ code snippet (Table 9)
    if len(doc.tables) > 9:
        table = doc.tables[9]
        for row in table.rows:
            for cell in row.cells:
                if "self.directed" in cell.text and "self.adj_list" in cell.text:
                    new_code = '''def __init__(self, nodes=None, edges=None, is_directed=False, weighted=True):
    self.nodes = list(nodes) if nodes else []
    self.is_directed = is_directed
    self.weighted = weighted
    self.adj = {}
    for node in self.nodes:
        self.adj[node] = []
    if edges:
        for edge in edges:
            u, v, w = edge if len(edge) == 3 else (edge[0], edge[1], 1)
            self.add_edge(u, v, w)'''
                    for para in cell.paragraphs:
                        if "self.directed" in para.text:
                            for run in para.runs:
                                if "self.directed" in run.text:
                                    run.text = new_code
                                    changes += 1
                                    print("  Fixed __init__ code")
                                    break

    # Fix the add_edge code snippet (Table 10)
    if len(doc.tables) > 10:
        table = doc.tables[10]
        for row in table.rows:
            for cell in row.cells:
                if "self.add_node" in cell.text or "weight" in cell.text:
                    new_code = '''def add_edge(self, u, v, w=1):
    self.adj[u].append((v, w))
    if not self.is_directed:
        self.adj[v].append((u, w))'''
                    for para in cell.paragraphs:
                        if "add_node" in para.text or "weight" in para.text:
                            for run in para.runs:
                                if "add_node" in run.text or "weight" in run.text:
                                    run.text = new_code
                                    changes += 1
                                    print("  Fixed add_edge code")
                                    break

    # Fix the adjacency matrix code snippet (Table 11)
    if len(doc.tables) > 11:
        table = doc.tables[11]
        for row in table.rows:
            for cell in row.cells:
                if "inf" in cell.text and "self.weighted" in cell.text:
                    new_code = '''n = len(self.nodes)
idx = {node: i for i, node in enumerate(self.nodes)}
matrix = [[0]*n for _ in range(n)]
for u in self.nodes:
    for v, w in self.adj[u]:
        matrix[idx[u]][idx[v]] = w'''
                    for para in cell.paragraphs:
                        if "inf" in para.text:
                            for run in para.runs:
                                if "inf" in run.text:
                                    run.text = new_code
                                    changes += 1
                                    print("  Fixed adjacency matrix code")
                                    break

    # Fix the edge list code snippet (Table 12)
    if len(doc.tables) > 12:
        table = doc.tables[12]
        for row in table.rows:
            for cell in row.cells:
                if "edge_key" in cell.text or "min(u, v)" in cell.text:
                    new_code = '''edges, seen = [], set()
for u in self.nodes:
    for v, w in self.adj[u]:
        if self.is_directed:
            edges.append((u, v, w))
        else:
            key = (min(u, v), max(u, v))
            if key not in seen:
                seen.add(key)
                edges.append((u, v, w))'''
                    for para in cell.paragraphs:
                        if "edge_key" in para.text or "min(u, v)" in para.text:
                            for run in para.runs:
                                if "edge_key" in run.text or "min(u, v)" in run.text:
                                    run.text = new_code
                                    changes += 1
                                    print("  Fixed edge list code")
                                    break

    # Fix the file write code snippet (Table 13)
    if len(doc.tables) > 13:
        table = doc.tables[13]
        for row in table.rows:
            for cell in row.cells:
                if "directed" in cell.text and "f.write" in cell.text:
                    new_code = '''f.write("directed\\n" if self.is_directed else "undirected\\n")
f.write("weighted\\n" if self.weighted else "unweighted\\n")
f.write(f"{len(self.nodes)}\\n")
f.write(f"{len(edges)}\\n")
f.write(" ".join(map(str, self.nodes)) + "\\n")
for u, v, w in edges:
    f.write(f"{u} {v} {w}\\n")'''
                    for para in cell.paragraphs:
                        if "directed" in para.text and "f.write" in para.text:
                            for run in para.runs:
                                if "directed" in run.text and "f.write" in run.text:
                                    run.text = new_code
                                    changes += 1
                                    print("  Fixed file write code")
                                    break

    # Fix the draw code snippet (Table 14)
    if len(doc.tables) > 14:
        table = doc.tables[14]
        for row in table.rows:
            for cell in row.cells:
                if "nx.DiGraph" in cell.text and "spring_layout" in cell.text:
                    new_code = '''import networkx as nx
G = nx.DiGraph() if self.is_directed else nx.Graph()
pos = nx.spring_layout(G, seed=42)
if self.is_directed:
    nx.draw_networkx_edges(G, pos, arrows=True, arrowsize=20)'''
                    for para in cell.paragraphs:
                        if "nx.DiGraph" in para.text:
                            for run in para.runs:
                                if "nx.DiGraph" in run.text:
                                    run.text = new_code
                                    changes += 1
                                    print("  Fixed draw code")
                                    break

    # ============================================================
    # FIX 3: Sua trong chuoi trich dan code (Table 8 - full code)
    # ============================================================
    if len(doc.tables) > 8:
        table = doc.tables[8]
        for row in table.rows:
            for cell in row.cells:
                if "graph_module" in cell.text or "self.adj_list" in cell.text:
                    # This is the full code table - needs complete replacement
                    new_code = '''"""
graph.py - Lop Do Thi co ban
Mo ta: Luu tru va quan ly do thi voi cac phep bieu dien:
  - Danh sach ke (Adjacency List)
  - Ma tran ke (Adjacency Matrix)
  - Danh sach canh (Edge List)
"""

import math


class Graph:
    """Lop Do Thi ho tro co huong/vo huong, co trong so/khong trong so"""

    def __init__(self, nodes=None, edges=None, is_directed=False, weighted=True):
        self.nodes = list(nodes) if nodes else []
        self.is_directed = is_directed
        self.weighted = weighted
        self.adj = {}
        for node in self.nodes:
            self.adj[node] = []
        if edges:
            for edge in edges:
                if len(edge) == 3:
                    u, v, w = edge
                else:
                    u, v = edge
                    w = 1
                self.add_edge(u, v, w)

    def add_edge(self, u, v, w=1):
        self.adj[u].append((v, w))
        if not self.is_directed:
            self.adj[v].append((u, w))

    def get_neighbors(self, node):
        return list(self.adj.get(node, []))

    def get_all_edges(self):
        edges, seen = [], set()
        for u in self.nodes:
            for v, w in self.adj[u]:
                if self.is_directed:
                    edges.append((u, v, w))
                else:
                    key = (min(u, v), max(u, v))
                    if key not in seen:
                        seen.add(key)
                        edges.append((u, v, w))
        return edges

    def get_adjacency_matrix(self):
        n = len(self.nodes)
        idx = {node: i for i, node in enumerate(self.nodes)}
        matrix = [[0]*n for _ in range(n)]
        for u in self.nodes:
            for v, w in self.adj[u]:
                matrix[idx[u]][idx[v]] = w
        return matrix, self.nodes

    def get_adjacency_list(self):
        return {node: list(self.adj[node]) for node in self.nodes}

    def get_edge_list(self):
        return self.get_all_edges()

    def draw_graph(self, filename="graph.png", title="Do Thi",
                   node_colors=None, highlight_edges=None):
        # ... (ve do thi bang networkx/matplotlib)
        pass

    def copy(self):
        edges = self.get_all_edges()
        return Graph(
            nodes=list(self.nodes), edges=list(edges),
            is_directed=self.is_directed, weighted=self.weighted,
        )'''
                    for para in cell.paragraphs:
                        if "graph_module" in para.text or "self.adj_list" in para.text:
                            for run in para.runs:
                                if "graph_module" in run.text or "self.adj_list" in run.text:
                                    run.text = new_code
                                    changes += 1
                                    print("  Fixed full code table")
                                    break

    # ============================================================
    # FIX 4: Sua trong cac doan van mo ta
    # ============================================================

    # Fix descriptions of class attributes
    attr_fixes = [
        ("directed là", "is_directed là"),
        ("adj_list lưu", "adj lưu"),
        ("adj_list chứa", "adj chứa"),
        ("adj_list dạng", "adj dạng"),
        ("get_nodes() trả về", "get_all_edges() trả về"),
        ("get_nodes() để", "get_all_edges() để"),
        ("ba phương thức công khai", "hai phương thức công khai"),
        ("get_neighbors() và get_nodes()", "get_neighbors() và get_all_edges()"),
    ]

    for para in doc.paragraphs:
        for old, new in attr_fixes:
            if old in para.text:
                if replace_in_paragraph(para, old, new):
                    changes += 1
                    print(f"  Fixed: '{old}' -> '{new}'")

    # ============================================================
    # FIX 5: Sua ten file trong cac tham chieu
    # ============================================================
    file_fixes = [
        ("graph_module.py", "graph.py"),
        ("graph_traversal.py", "q3_traversal.py"),
        ("graph_shortest_path.py", "q5_shortest.py"),
        ("graph_advanced_algorithms.py", "q73q74_mst.py"),
    ]

    for para in doc.paragraphs:
        for old, new in file_fixes:
            if old in para.text:
                if replace_in_paragraph(para, old, new):
                    changes += 1
                    print(f"  Fixed file: '{old}' -> '{new}'")

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for old, new in file_fixes:
                        if old in para.text:
                            if replace_in_paragraph(para, old, new):
                                changes += 1
                                print(f"  Fixed file in table: '{old}' -> '{new}'")

    # Save
    print(f"\nTong so thay doi: {changes}")
    doc.save(OUTPUT)
    print(f"Da luu: {OUTPUT}")


if __name__ == "__main__":
    main()
